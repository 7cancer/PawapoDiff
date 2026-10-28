"""Read-only OOXML inspection; geometry is expressed in slide points."""

from io import BytesIO
from hashlib import sha256
import logging
import zlib
from zipfile import ZipFile, BadZipFile
from lxml import etree
from pptx import Presentation
from pptx.shapes.picture import Picture
from .limits import MAX_EXPANDED_BYTES, MAX_XML_PART_BYTES, MAX_XML_TOTAL_BYTES
from .errors import InputError
from .media import plan_media, target_part
from .translations import normalize, translation_table

EMU = 12700


def validate(path):
    part = None
    try:
        with ZipFile(path) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_EXPANDED_BYTES:
                raise InputError(
                    "展開後のサイズが60GBを超えています。",
                    code="EXPANDED_SIZE_LIMIT",
                    actions=["画像や動画を圧縮するか、資料を分割してください。"],
                )
            xml_parts = [
                i for i in z.infolist() if i.filename.endswith((".xml", ".rels"))
            ]
            if (
                any(i.file_size > MAX_XML_PART_BYTES for i in xml_parts)
                or sum(i.file_size for i in xml_parts) > MAX_XML_TOTAL_BYTES
            ):
                raise InputError(
                    "XML構造情報が上限（1パーツ64MB・合計512MB）を超えています。",
                    code="XML_SIZE_LIMIT",
                    actions=[
                        "複雑な図形や表が多い資料は、資料を分割してお試しください。"
                    ],
                )
            names = z.namelist()
            name_set = set(names)
            if len(names) != len(set(names)):
                raise InputError(
                    "内部ファイル名が重複したPPTXです。",
                    code="DUPLICATE_PART",
                    actions=["PowerPointで別名保存した資料を指定してください。"],
                )
            if any(i.flag_bits & 1 for i in z.infolist()):
                raise InputError(
                    "暗号化されたZIPには対応していません。", code="ENCRYPTED_ZIP"
                )
            required = {"ppt/presentation.xml", "[Content_Types].xml", "_rels/.rels"}
            if not required <= set(names):
                raise InputError(
                    "PPTXに必要な内部ファイルがありません。拡張子だけがpptxの資料や不完全なダウンロードの可能性があります。",
                    code="NOT_PPTX",
                    actions=[
                        "PowerPointから通常の「PowerPointプレゼンテーション (*.pptx)」として保存したファイルを選択してください。"
                    ],
                )
            part = "ppt/presentation.xml"
            presentation = etree.fromstring(
                z.read(part), etree.XMLParser(resolve_entities=False, no_network=True)
            )
            if (
                etree.QName(presentation).namespace
                == "http://purl.oclc.org/ooxml/presentationml/main"
            ):
                raise InputError(
                    "Strict Open XML形式のPPTXには対応していません。通常のPPTXとは内部形式が異なります。",
                    code="STRICT_OOXML",
                    part=part,
                    actions=[
                        "PowerPointの「名前を付けて保存」で「Strict」ではない「PowerPointプレゼンテーション (*.pptx)」を選択してください。"
                    ],
                )
            media_plan = plan_media(z)
            for name in names:
                part = name
                if "vbaproject" in name.lower() or "activex" in name.lower():
                    raise InputError(
                        "マクロ・ActiveXを含むファイルには対応していません。",
                        code="ACTIVE_CONTENT",
                        part=name,
                    )
                if name in media_plan.removed:
                    continue
                if name.endswith(".rels"):
                    root = etree.fromstring(
                        media_plan.replacements[name]
                        if name in media_plan.replacements
                        else z.read(name),
                        etree.XMLParser(resolve_entities=False, no_network=True),
                    )
                    # External media/OLE are not opened. Hyperlinks are safe to render without following.
                    for r in root:
                        if not isinstance(r.tag, str):
                            continue
                        kind = r.get("Type", "").rsplit("/", 1)[-1]
                        if r.get("TargetMode") == "External":
                            if kind != "hyperlink":
                                raise InputError(
                                    f"外部参照（{kind}）を含むため、安全に読み込めません。",
                                    code="EXTERNAL_REFERENCE",
                                    part=name,
                                    actions=[
                                        "PowerPointの「ファイル → 情報 → ファイルへのリンクの編集」でリンク先のExcel・画像等を確認し、埋め込みまたはリンク解除したコピーを保存してください。外部リンクは更新しません。"
                                    ],
                                )
                        else:
                            target = target_part(name, r.get("Target", ""))
                            if target not in name_set:
                                raise InputError(
                                    "内部の参照先が欠落しています。",
                                    code="MISSING_PART",
                                    part=name,
                                    technical=f"Relationship {r.get('Id')} -> {target}",
                                    actions=[
                                        "PowerPointの修復で開けるか確認し、修復後のコピーを保存してください。"
                                    ],
                                )
            return media_plan
    except (BadZipFile, zlib.error, EOFError) as e:
        raise InputError(
            "PPTXの圧縮データを読み取れません。破損、暗号化・権限保護、旧形式PPTの可能性があります。",
            code="INVALID_PACKAGE",
            part=part,
            actions=[
                "PowerPointで直接開き、パスワード・権限保護の有無を確認してください。保護を解除できる権限がある場合は、通常のPPTXとしてコピーを保存してください。",
                "OneDrive等のファイルはダウンロードを完了してください。拡張子の変更だけでは変換できません。",
            ],
        ) from e
    except etree.XMLSyntaxError as e:
        raise InputError(
            "PowerPoint内部のXMLが破損しているか、解析上限を超えています。",
            code="INVALID_XML",
            part=part,
            technical=f"{type(e).__name__}: line {e.lineno}, column {e.position[1]}",
        ) from e


def color_value(color):
    if color.type is None:
        return None
    # Include system/preset/HSL/scRGB colors and tint/alpha transforms. Using
    # theme_color alone collapses all non-theme colors to NOT_THEME_COLOR.
    return format_xml(color._xFill)


def format_xml(element):
    """Keep formatting and geometry, ignoring text content and proofing metadata."""
    if element is None:
        return ""
    copy = etree.fromstring(etree.tostring(element))
    for node in copy.iter():
        if not isinstance(node.tag, str):
            continue  # XML comments and processing instructions are legal.
        if etree.QName(node).localname in ("t", "fld"):
            node.text = ""
        for attr in list(node.attrib):
            if etree.QName(attr).localname in (
                "lang",
                "altLang",
                "dirty",
                "err",
                "smtClean",
                "smtId",
            ):
                del node.attrib[attr]
    return etree.tostring(copy).decode()


def inherited_signature(slide):
    parts = [
        format_xml(slide.slide_layout._element),
        format_xml(slide.slide_layout.slide_master._element),
    ]
    for rel in slide.slide_layout.slide_master.part.rels.values():
        if not rel.is_external and rel.reltype.endswith("/theme"):
            parts.append(sha256(rel.target_part.blob).hexdigest())
    return sha256("".join(parts).encode()).hexdigest()


def text_style(frame):
    paragraphs = []
    for p in frame.paragraphs:
        runs = []
        for r in p.runs:
            if not r.text.strip():
                continue
            f = r.font
            style = (
                f.name,
                f.size.pt if f.size else None,
                f.bold,
                f.italic,
                color_value(f.color),
            )
            # Compare transitions, not translated run lengths or word counts.
            if not runs or runs[-1] != style:
                runs.append(style)
        paragraphs.append(
            {
                "alignment": str(p.alignment),
                "level": p.level,
                "runs": runs,
                "defaults": format_xml(p._p.pPr),
            }
        )
    # Paragraph count may naturally change on translation; preserve distinct style patterns.
    distinct = []
    for p in paragraphs:
        if p not in distinct:
            distinct.append(p)
    return {
        "patterns": distinct,
        "anchor": str(frame.vertical_anchor),
        "margins": [
            getattr(frame, "margin_" + side) / EMU
            for side in ("left", "top", "right", "bottom")
        ],
        "wrap": frame.word_wrap,
        "body": format_xml(frame._txBody.bodyPr),
    }


def inspect(path, *, progress=None):
    media_plan = validate(path)
    source_media = None
    page = 0
    part = "ppt/presentation.xml"
    shape_id = None
    try:
        # python-pptx eagerly loads every binary part. Give it a metadata-only
        # package, and hash the original media in bounded chunks when needed.
        metadata = BytesIO()
        with ZipFile(path) as source, ZipFile(metadata, "w") as target:
            for info in source.infolist():
                if info.filename in media_plan.removed:
                    continue
                part = info.filename
                data = (
                    (
                        media_plan.replacements[info.filename]
                        if info.filename in media_plan.replacements
                        else source.read(info.filename)
                    )
                    if info.filename.endswith((".xml", ".rels"))
                    else b""
                )
                if data:
                    # Identify a broken XML part before python-pptx loses its name.
                    try:
                        root = etree.fromstring(
                            data,
                            etree.XMLParser(resolve_entities=False, no_network=True),
                        )
                    except etree.XMLSyntaxError as error:
                        raise InputError(
                            "PowerPoint内部のXMLが破損しているか、解析上限を超えています。",
                            code="INVALID_XML",
                            part=part,
                            technical=f"{type(error).__name__}: line {error.lineno}, column {error.position[1]}",
                        ) from error
                    if root.getroottree().docinfo.doctype:
                        raise InputError(
                            "DTDを含むXMLは安全のため読み込みません。",
                            code="XML_DTD",
                            part=part,
                        )
                target.writestr(info.filename, data)
        metadata.seek(0)
        part = "ppt/presentation.xml"
        prs = Presentation(metadata)
        source_media = ZipFile(path)
        image_hashes = {}
        slides = []
        inherited = []
        warnings = []
        signature_cache = {}
        recoverable = (
            AttributeError,
            ValueError,
            TypeError,
            KeyError,
            IndexError,
            ZeroDivisionError,
            NotImplementedError,
        )

        def warning(key, field, error):
            warnings.append(
                {
                    "page": page,
                    "shape_id": str(key),
                    "field": field,
                    "message": "この要素の" + field + "は画像で確認してください。",
                    "technical": f"{type(error).__name__}: {error}"[:500],
                }
            )
            logging.getLogger(__name__).warning(
                "Partial analysis: page=%s part=%s shape_id=%s field=%s exception=%s",
                page,
                part,
                key,
                field,
                type(error).__name__,
            )

        def optional(key, field, action, fallback=None):
            try:
                return action()
            except recoverable as error:
                warning(key, field, error)
                return fallback

        def walk(shapes, prefix="", transform=(1, 1, 0, 0)):
            nonlocal shape_id
            result = []
            for index, sh in enumerate(shapes):
                shape_id = optional(
                    prefix + str(index),
                    "図形ID",
                    lambda: sh.shape_id,
                    f"unknown-{index}",
                )
                key = f"{prefix}{shape_id}"
                before = len(warnings)
                try:
                    items = walk_one([sh], prefix, transform)
                except recoverable as error:
                    warning(key, "構造・書式・配置", error)
                    # Preserve available text; never silently omit the shape.
                    text = "\n".join(sh._element.xpath(".//a:t/text()"))
                    items = [
                        {
                            "id": key,
                            "name": "解析要確認の要素 " + key,
                            "kind": "unsupported",
                            "box": [
                                0,
                                0,
                                prs.slide_width / EMU,
                                prs.slide_height / EMU,
                            ],
                            "rotation": 0,
                            "text": text,
                            "style": None,
                            "image": None,
                            "geometry_unknown": True,
                        }
                    ]
                if len(warnings) > before:
                    for item in items:
                        item["analysis_incomplete"] = True
                result.extend(items)
            return result

        def walk_one(shapes, prefix="", transform=(1, 1, 0, 0)):
            result = []
            sx, sy, ox, oy = transform
            for sh in shapes:
                key = f"{prefix}{sh.shape_id}"
                dimensions = [sh.left, sh.top, sh.width, sh.height]
                geometry_unknown = any(value is None for value in dimensions)
                if geometry_unknown:
                    warning(
                        key,
                        "配置・寸法",
                        ValueError("Geometry unavailable in shape/layout/master"),
                    )
                    box = [0, 0, prs.slide_width / EMU, prs.slide_height / EMU]
                else:
                    box = [
                        dimensions[0] / EMU * sx + ox,
                        dimensions[1] / EMU * sy + oy,
                        dimensions[2] / EMU * sx,
                        dimensions[3] / EMU * sy,
                    ]
                if sh.shape_type == 6:
                    x = sh._element.grpSpPr.xfrm
                    if (
                        x is not None
                        and x.chExt is not None
                        and x.chOff is not None
                        and not geometry_unknown
                    ):
                        gx = sh.width / x.chExt.cx if x.chExt.cx else 1
                        gy = sh.height / x.chExt.cy if x.chExt.cy else 1
                        child_transform = (
                            sx * gx,
                            sy * gy,
                            box[0] - x.chOff.x / EMU * sx * gx,
                            box[1] - x.chOff.y / EMU * sy * gy,
                        )
                    else:
                        warning(
                            key,
                            "グループ内の配置",
                            ValueError("Group transform unavailable"),
                        )
                        child_transform = transform
                        geometry_unknown = True
                    result.append(
                        {
                            "id": key,
                            "kind": "group",
                            "geometry_unknown": geometry_unknown,
                            "name": sh.name,
                            "box": box,
                            "rotation": sh.rotation,
                            "text": "",
                            "style": None,
                            "image": None,
                        }
                    )
                    children = walk(sh.shapes, key + "/", child_transform)
                    if geometry_unknown:
                        for child in children:
                            child["geometry_unknown"] = True
                    result.extend(children)
                    continue
                item = {
                    "id": key,
                    "kind": str(sh.shape_type),
                    "geometry_unknown": geometry_unknown,
                    "name": sh.name,
                    "box": box,
                    "rotation": sh.rotation,
                    "text": "",
                    "style": None,
                    "image": None,
                }
                if sh.has_text_frame:
                    item.update(
                        text=sh.text,
                        style=optional(
                            key, "文字書式", lambda: text_style(sh.text_frame)
                        ),
                    )
                if isinstance(sh, Picture):
                    rid = sh._pic.blip_rId
                    name = str(sh.part.related_part(rid).partname).lstrip("/")
                    if name not in image_hashes:
                        digest = sha256()
                        with source_media.open(name) as stream:
                            for chunk in iter(lambda: stream.read(1024 * 1024), b""):
                                digest.update(chunk)
                        image_hashes[name] = digest.hexdigest()
                    item["image"] = image_hashes[name]
                result.append(item)
                if sh.has_table:
                    table = sh.table
                    y = box[1]
                    for row_i, row in enumerate(table.rows):
                        x = box[0]
                        for col_i, cell in enumerate(row.cells):
                            w = table.columns[col_i].width / EMU * sx
                            if not cell.is_spanned:
                                cw = (
                                    sum(
                                        table.columns[c].width
                                        for c in range(col_i, col_i + cell.span_width)
                                    )
                                    / EMU
                                    * sx
                                )
                                ch = (
                                    sum(
                                        table.rows[r].height
                                        for r in range(row_i, row_i + cell.span_height)
                                    )
                                    / EMU
                                    * sy
                                )
                                result.append(
                                    {
                                        "id": f"{key}/cell/{row_i}/{col_i}",
                                        "kind": "cell",
                                        "geometry_unknown": geometry_unknown,
                                        "name": f"{sh.name} ({row_i + 1},{col_i + 1})",
                                        "box": [x, y, cw, ch],
                                        "rotation": 0,
                                        "text": cell.text,
                                        "style": optional(
                                            f"{key}/cell/{row_i}/{col_i}",
                                            "セルの文字書式",
                                            lambda: text_style(cell.text_frame),
                                        ),
                                        "image": None,
                                    }
                                )
                            x += w
                        y += row.height / EMU * sy
            return result

        if not prs.slides:
            raise InputError("スライドが1ページもありません。", code="NO_SLIDES")
        if (
            not prs.slide_width
            or not prs.slide_height
            or min(prs.slide_width, prs.slide_height) <= 0
        ):
            raise InputError("スライド寸法が不正です。", code="INVALID_SLIDE_SIZE")
        for page, slide in enumerate(prs.slides, 1):
            part = str(slide.part.partname)
            if progress:
                progress(page, len(prs.slides))
            slides.append(walk(slide.shapes))
            key = str(slide.slide_layout.part.partname)
            if key not in signature_cache:
                signature_cache[key] = optional(
                    "inherited",
                    "テーマ・マスター・レイアウト",
                    lambda: inherited_signature(slide),
                )
            inherited.append(signature_cache[key])
            if signature_cache[key] is None:
                warning(
                    "inherited",
                    "テーマ・マスター・レイアウト",
                    ValueError("Inherited formatting unavailable"),
                )
        return {
            "width": prs.slide_width / EMU,
            "height": prs.slide_height / EMU,
            "slides": slides,
            "inherited": inherited,
            "warnings": warnings,
            "notices": media_plan.notices,
        }
    except InputError:
        raise
    except (OSError, MemoryError):
        raise
    except Exception as e:
        logging.getLogger(__name__).exception(
            "PPTX analysis failed: page=%s part=%s shape_id=%s", page, part, shape_id
        )
        raise InputError(
            "PPTXの内部構造を解析できませんでした。拡張子の問題とは限りません。",
            code="ANALYSIS_FAILED",
            page=page or None,
            part=part,
            shape_id=shape_id,
            technical=f"{type(e).__name__}: {e}"[:1500],
            actions=[
                "下の技術情報で、失敗した内部ファイル・ページ・図形を確認してください。",
                "診断ログを保存してください。PowerPointで手動で開ける場合は、その旨もお知らせください。",
            ],
        ) from e
    finally:
        if source_media is not None:
            source_media.close()


def pair_shapes(left, right):
    available = {b["id"]: b for b in right}
    pairs, remaining = [], []
    for a in left:
        b = available.get(a["id"])
        if b and a["kind"] == b["kind"]:
            pairs.append((a, available.pop(a["id"]), False))
        else:
            remaining.append(a)
    for a in remaining:
        candidates = [b for b in available.values() if b["kind"] == a["kind"]]
        b = min(
            candidates,
            key=lambda b: sum(abs(x - y) for x, y in zip(a["box"], b["box"])),
            default=None,
        )
        if b:
            available.pop(b["id"])
        pairs.append((a, b, True))
    pairs.extend((None, b, True) for b in available.values())
    return pairs


def compare(left, right, threshold=1):
    if len(left["slides"]) != len(right["slides"]):
        raise InputError(
            f"ページ数が一致しません（原本{len(left['slides'])}ページ、英訳版{len(right['slides'])}ページ）。ページ数が同じ資料を指定してください。",
            code="PAGE_COUNT_MISMATCH",
        )
    if (left["width"], left["height"]) != (right["width"], right["height"]):
        raise InputError(
            f"スライドの寸法が一致しません（原本{left['width']:g}×{left['height']:g}pt、英訳版{right['width']:g}×{right['height']:g}pt）。",
            code="SLIDE_SIZE_MISMATCH",
        )
    pages, text_pairs = [], []
    for page, (aa, bb) in enumerate(zip(left["slides"], right["slides"]), 1):
        changes = []
        inherited_left = left.get("inherited", [None] * len(left["slides"]))[page - 1]
        inherited_right = right.get("inherited", [None] * len(right["slides"]))[
            page - 1
        ]
        inherited_unknown = any(
            warning["page"] == page and warning["shape_id"] == "inherited"
            for doc in (left, right)
            for warning in doc.get("warnings", [])
        )
        if inherited_left != inherited_right or inherited_unknown:
            marker = {
                "id": "inherited",
                "name": "テーマ・マスター・レイアウト",
                "box": [0, 0, left["width"], left["height"]],
            }
            changes.append(
                {
                    "kind": "structure",
                    "highlight": "other",
                    "fields": [
                        "継承された書式・レイアウトの解析が不完全（画像で要確認）"
                        if inherited_unknown
                        else "継承された書式・レイアウトの変更（要確認）"
                    ],
                    "left": marker,
                    "right": marker,
                    "uncertain": True,
                }
            )
        for a, b, uncertain in pair_shapes(aa, bb):
            source = a.get("text", "") if a else ""
            target = b.get("text", "") if b else ""
            source_key, target_key = normalize(source), normalize(target)
            fields = []
            if (a and a.get("analysis_incomplete")) or (
                b and b.get("analysis_incomplete")
            ):
                uncertain = True
                fields.append("解析できない項目あり（画像で要確認）")
            delta = None
            layout = formatting = False
            if not a or not b:
                fields.append("要素の追加" if b else "要素の削除")
            else:
                delta = [round(y - x, 4) for x, y in zip(a["box"], b["box"])]
                for i, label in enumerate(
                    ["横位置", "縦位置", "幅", "高さ"]
                    if not (a.get("geometry_unknown") or b.get("geometry_unknown"))
                    else []
                ):
                    if abs(delta[i]) > 0.00001 and abs(delta[i]) >= threshold:
                        fields.append(f"{label}: {delta[i]:+.2f} pt")
                        layout = True
                if abs(a["rotation"] - b["rotation"]) > 0.001:
                    fields.append("回転")
                    layout = True
                style_unknown = any(
                    item.get("analysis_incomplete") and item["style"] is None
                    for item in (a, b)
                )
                if not style_unknown and a["style"] != b["style"]:
                    fields.append("文字書式・配置")
                    formatting = True
                if a["image"] != b["image"]:
                    fields.append("画像")
                if source_key != target_key:
                    changes.append(
                        {
                            "kind": "text",
                            "fields": ["文字内容（翻訳を含む）"],
                            "left": a,
                            "right": b,
                            "uncertain": uncertain,
                        }
                    )
            if fields or uncertain:
                highlight = (
                    "mixed"
                    if layout and formatting
                    else "layout"
                    if layout
                    else "format"
                    if formatting
                    else "other"
                )
                changes.append(
                    {
                        "kind": "structure",
                        "highlight": highlight,
                        "fields": fields or ["対応付け要確認"],
                        "delta": delta,
                        "left": a,
                        "right": b,
                        "uncertain": uncertain,
                    }
                )
            if source_key or target_key:
                notes = []
                if uncertain:
                    notes.append("対応付け要確認")
                if (a and a.get("analysis_incomplete")) or (
                    b and b.get("analysis_incomplete")
                ):
                    notes.append("解析が一部不完全・画像で要確認")
                text_pairs.append(
                    {
                        "page": page,
                        "shape": (a or b)["name"],
                        "source": source,
                        "target": target,
                        "uncertain": uncertain,
                        "notes": notes,
                    }
                )
        pages.append({"page": page, "changes": changes})
    rows, issues = translation_table(text_pairs)
    return {
        "width": left["width"],
        "height": left["height"],
        "pages": pages,
        "issues": issues,
        "translations": rows,
        "threshold": threshold,
        "notices": [
            dict(n, side=side)
            for side, doc in (("left", left), ("right", right))
            for n in doc.get("notices", [])
        ],
        "warnings": [
            dict(w, side=side)
            for side, doc in (("left", left), ("right", right))
            for w in doc.get("warnings", [])
        ],
    }
