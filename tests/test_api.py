from fastapi.testclient import TestClient
from pptx import Presentation
from PIL import Image
import time
from pawapodiff.server import create_app


def test_job_report_and_security(tmp_path):
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[6])
    s.shapes.add_textbox(0, 0, 100000, 100000).text = "日本語"
    source = tmp_path / "input.pptx"
    p.save(source)

    def renderer(path, dest):
        dest.mkdir()
        out = dest / "0001.png"
        Image.new("RGB", (160, 90), "white").save(out)
        return [out]

    app = create_app("test-token", tmp_path / "work", renderer)
    with TestClient(app, base_url="http://127.0.0.1") as c:
        assert c.get("/api/jobs/unknown").status_code == 403
        h = {"X-Pawapo-Token": "test-token"}
        assert (
            c.get(
                "/api/jobs/unknown", headers={**h, "Origin": "http://evil.example"}
            ).status_code
            == 403
        )
        assert (
            c.get(
                "/api/jobs/unknown", headers={**h, "Host": "evil.example"}
            ).status_code
            == 403
        )
        data = source.read_bytes()
        response = c.post(
            "/api/jobs",
            headers=h,
            files={"left": ("a.pptx", data), "right": ("b.pptx", data)},
        )
        assert response.status_code == 200, response.text
        key = response.json()["id"]
        for _ in range(100):
            status = c.get(f"/api/jobs/{key}", headers=h).json()
            if status["status"] in ("done", "error"):
                break
            time.sleep(0.02)
        assert status["status"] == "done", status
        assert status["result"]["issues"][0]["kind"] == "日本語の残存候補"
        assert (
            c.get(f"/api/jobs/{key}/images/left/1", headers=h).headers["content-type"]
            == "image/png"
        )
        assert c.get(f"/api/jobs/{key}/images/diff/1", headers=h).status_code == 200
        assert (
            c.post(
                f"/api/jobs/{key}/threshold", headers=h, json={"threshold": 0.1}
            ).status_code
            == 200
        )
        assert (
            c.post(
                f"/api/jobs/{key}/threshold", headers=h, json={"threshold": -1}
            ).status_code
            == 422
        )
        assert (
            b"data:image/png;base64,"
            in c.get(f"/api/jobs/{key}/report", headers=h).content
        )
        assert status["result"]["translations"][0]["source"] == "日本語"
        response = c.post("/api/jobs", headers=h, files={"left": ("a.pptx", data)})
        assert response.status_code == 422
        assert c.post(f"/api/jobs/{key}/export", headers=h, json={}).status_code in (
            404,
            405,
        )


def test_diagnostics_download_requires_local_authentication(tmp_path):
    app = create_app("diagnostic-token", tmp_path / "work")
    with TestClient(app, base_url="http://127.0.0.1") as client:
        assert client.get("/api/diagnostics").status_code == 403
        response = client.get(
            "/api/diagnostics", headers={"X-Pawapo-Token": "diagnostic-token"}
        )
        assert response.status_code == 200
        assert "PawapoDiff-diagnostics.zip" in response.headers["content-disposition"]
        from zipfile import ZipFile
        import io

        with ZipFile(io.BytesIO(response.content)) as archive:
            assert archive.namelist()
            assert all(
                b"diagnostic-token" not in archive.read(name)
                for name in archive.namelist()
            )
