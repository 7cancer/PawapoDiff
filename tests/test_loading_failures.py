"""Real OOXML edge cases and end-to-end failures, including a 200-page pair."""

from io import BytesIO
import json
import time
from zipfile import ZipFile

from fastapi.testclient import TestClient
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches
import pytest

import pawapodiff.engine as engine
from pawapodiff.errors import describe, office_diagnostic
from pawapodiff.render import RenderError, _run_worker
from pawapodiff.server import create_app


def deck(path, pages=1):
    p = Presentation()
    for i in range(pages):
        slide = p.slides.add_slide(p.slide_layouts[6])
        slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(2), Inches(1)
        ).text = f"社外秘本文 {i + 1}"
    p.save(path)
    return p


def rewrite(source, target, replacements):
    with ZipFile(source) as a, ZipFile(target, "w") as b:
        for item in a.infolist():
            if (
                replacements.get(item.filename) is None
                and item.filename in replacements
            ):
                continue
            b.writestr(item, replacements.get(item.filename, a.read(item.filename)))


def test_missing_geometry_preserves_text_without_inventing_layout_differences(tmp_path):
    source = tmp_path / "source.pptx"
    p = deck(source)
    moved = p.slides[0].shapes.add_textbox(
        Inches(4), Inches(1), Inches(2), Inches(1)
    )
    moved.text = "配置の変更は表示する文章"
    p.save(source)
    normal = engine.inspect(source)
    shape = p.slides[0].shapes[0]
    shape._element.spPr.remove(shape._element.spPr.xfrm)
    moved.left += Inches(0.5)
    p.save(source)
    partial = engine.inspect(source)
    assert partial["slides"][0][0]["text"] == normal["slides"][0][0]["text"]
    assert partial["warnings"][0]["page"] == 1
    result = engine.compare(normal, partial)
    assert result["warnings"] and result["translations"][0]["uncertain"]
    changes = result["pages"][0]["changes"]
    assert len(changes) == 1
    assert changes[0]["left"]["id"] == str(moved.shape_id)
    assert changes[0]["highlight"] == "layout"
    assert all(
        "解析できない" not in field for change in changes for field in change["fields"]
    )


def test_legal_xml_comments_and_system_colors_are_supported(tmp_path):
    source = tmp_path / "source.pptx"
    p = deck(source)
    run = p.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    fill = OxmlElement("a:solidFill")
    color = OxmlElement("a:sysClr")
    color.set("val", "windowText")
    color.set("lastClr", "000000")
    fill.append(color)
    run._r.get_or_add_rPr().append(fill)
    p.slide_layouts[6]._element.append(etree.Comment("valid XML comment"))
    p.save(source)
    original = engine.inspect(source)
    assert not original["warnings"]
    color.set("lastClr", "FF0000")
    p.save(source)
    result = engine.compare(original, engine.inspect(source))
    assert any(
        field.startswith("文字色")
        for c in result["pages"][0]["changes"]
        for field in c["fields"]
    )


def test_bad_optional_format_does_not_drop_text_or_other_slides(tmp_path):
    source = tmp_path / "source.pptx"
    p = deck(source, 2)
    paragraph = p.slides[0].shapes[0].text_frame.paragraphs[0]
    paragraph._p.get_or_add_pPr().set("algn", "futureAlignment")
    p.save(source)
    result = engine.inspect(source)
    assert len(result["slides"]) == 2
    assert result["slides"][0][0]["text"] == "社外秘本文 1"
    assert result["slides"][0][0]["analysis_incomplete"]
    assert result["warnings"][0]["page"] == 1
    assert not result["slides"][1][0].get("analysis_incomplete")


def test_group_without_transform_keeps_individual_child_texts(tmp_path):
    source = tmp_path / "source.pptx"
    p = deck(source)
    group = p.slides[0].shapes.add_group_shape()
    for text in ("子要素1", "子要素2"):
        group.shapes.add_textbox(0, 0, Inches(1), Inches(1)).text = text
    group._element.grpSpPr.remove(group._element.grpSpPr.xfrm)
    p.save(source)
    result = engine.inspect(source)
    assert {"子要素1", "子要素2"} <= {s["text"] for s in result["slides"][0]}
    assert result["warnings"]


@pytest.mark.parametrize(
    "mode,code,part",
    [
        ("missing", "MISSING_PART", "ppt/_rels/presentation.xml.rels"),
        ("strict", "STRICT_OOXML", "ppt/presentation.xml"),
        ("xml", "INVALID_XML", "ppt/slides/slide1.xml"),
        ("external", "EXTERNAL_REFERENCE", "ppt/slides/_rels/slide1.xml.rels"),
    ],
)
def test_specific_package_errors_identify_part(tmp_path, mode, code, part):
    source, bad = tmp_path / "source.pptx", tmp_path / "bad.pptx"
    deck(source)
    with ZipFile(source) as z:
        data = z.read(part)
    changes = {}
    if mode == "missing":
        changes["ppt/slides/slide1.xml"] = None
    elif mode == "strict":
        changes[part] = data.replace(
            b"http://schemas.openxmlformats.org/presentationml/2006/main",
            b"http://purl.oclc.org/ooxml/presentationml/main",
        )
    elif mode == "xml":
        changes[part] = b"<broken"
    else:
        changes[part] = data.replace(
            b"</Relationships>",
            b'<Relationship Id="linked" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject" Target="file:///private-company-folder/data.xlsx" TargetMode="External"/></Relationships>',
        )
    rewrite(source, bad, changes)
    with pytest.raises(engine.InputError) as caught:
        engine.inspect(bad)
    diagnostic = describe(caught.value)
    assert diagnostic["code"] == code
    assert diagnostic["part"] == part
    assert "private-company-folder" not in json.dumps(diagnostic)


@pytest.mark.parametrize(
    "code,expected",
    [
        (0x80040154, "POWERPOINT_NOT_INSTALLED"),
        (0x8001010A, "POWERPOINT_BUSY"),
        (0x800706BA, "POWERPOINT_DISCONNECTED"),
        (0x8007000E, "OFFICE_MEMORY_EXHAUSTED"),
        (0x80070070, "DISK_FULL"),
        (0x80070005, "OFFICE_ACCESS_DENIED"),
    ],
)
def test_hresult_and_office_nested_codes_have_actionable_messages(code, expected):
    error = Exception("Office error")
    error.hresult = -2147352567  # DISP_E_EXCEPTION, inspect Office's inner scode.
    error.excepinfo = (0, "PowerPoint", "detail", None, 0, code - 2**32)
    data = office_diagnostic(error, "資料を読み取り専用で開く")
    assert data["code"] == expected and data["actions"]


def wait(client, key, headers):
    for _ in range(500):
        status = client.get("/api/jobs/" + key, headers=headers).json()
        if status["status"] in ("error", "done"):
            return status
        time.sleep(0.01)
    pytest.fail("Job did not complete")


def test_200_page_pair_comparison_and_layout_signature_cache(tmp_path, monkeypatch):
    source = tmp_path / "large.pptx"
    deck(source, 200)
    calls = []
    original = engine.inherited_signature

    def signature(slide):
        calls.append(str(slide.slide_layout.part.partname))
        return original(slide)

    monkeypatch.setattr(engine, "inherited_signature", signature)

    def renderer(path, destination):
        destination.mkdir()
        images = []
        for index in range(1, 201):
            out = destination / f"{index:04}.png"
            Image.new("RGB", (32, 18)).save(out)
            images.append(out)
        return images

    app = create_app("test", tmp_path / "work", renderer)
    headers = {"X-Pawapo-Token": "test"}
    with TestClient(app, base_url="http://127.0.0.1") as client:
        response = client.post(
            "/api/jobs",
            headers=headers,
            files={
                side: ("large.pptx", source.read_bytes()) for side in ("left", "right")
            },
        )
        status = wait(client, response.json()["id"], headers)
        assert status["status"] == "done", status
        assert status["pages"] == 200
        assert len(status["result"]["translations"]) == 200
        assert not status["result"]["warnings"]
        assert len(calls) == 2  # One layout per document, not 200 serializations.
        archive = client.get("/api/diagnostics", headers=headers)
        with ZipFile(BytesIO(archive.content)) as z:
            data = json.loads(z.read("diagnostic-summary.json"))
            job = next(iter(data["jobs"].values()))
            assert job["files"]["left"]["pages"] == 200
            assert any("200/200" in event["stage"] for event in job["events"])
            assert "社外秘本文" not in z.read("diagnostic-summary.txt").decode(
                "utf-8-sig"
            )


def test_input_error_is_logged_and_summary_is_never_empty(
    tmp_path, monkeypatch, caplog
):
    import pawapodiff.diagnostics as diagnostics

    monkeypatch.setattr(diagnostics, "diagnostic_files", lambda: [])
    headers = {"X-Pawapo-Token": "test"}
    app = create_app("test", tmp_path / "work")
    with TestClient(app, base_url="http://127.0.0.1") as client:
        response = client.post(
            "/api/jobs",
            headers=headers,
            files={side: ("broken.pptx", b"bad") for side in ("left", "right")},
        )
        status = wait(client, response.json()["id"], headers)
        assert status["status"] == "error"
        assert status["diagnostic"]["code"] == "INVALID_PACKAGE"
        assert status["diagnostic"]["document"] == "日本語原本"
        assert status["diagnostic"]["technical"]
        assert "Job failed" in caplog.text
        data = client.get("/api/diagnostics", headers=headers).content
        with ZipFile(BytesIO(data)) as z:
            summary = z.read("diagnostic-summary.txt").decode("utf-8-sig")
            assert (
                "INVALID_PACKAGE" in summary
                and "日本語原本" in summary
                and "0.4.3" in summary
            )
            assert len(z.read("diagnostic-summary.json")) > 100


def test_upload_errors_also_appear_in_diagnostic_summary(tmp_path):
    app = create_app("test", tmp_path / "work")
    headers = {"X-Pawapo-Token": "test"}
    with TestClient(app, base_url="http://127.0.0.1") as client:
        assert (
            client.post(
                "/api/jobs", headers=headers, files={"left": ("bad.ppt", b"bad")}
            ).status_code
            == 422
        )
        with ZipFile(
            BytesIO(client.get("/api/diagnostics", headers=headers).content)
        ) as z:
            summary = json.loads(z.read("diagnostic-summary.json"))
            assert (
                next(iter(summary["jobs"].values()))["diagnostic"]["code"]
                == "UPLOAD_HTTP_422"
            )


def test_open_timeout_can_allow_large_decks_without_disabling_watchdog(tmp_path):
    from tests.test_office_worker import command

    args, state, stderr = command(
        tmp_path,
        "write_state(state,{'status':'running','stage':'資料を読み取り専用で開く'});time.sleep(1.2);write_state(state,{'status':'done'})",
    )
    progress = []
    assert (
        _run_worker(
            args, state, stderr, timeout=1, open_timeout=3, progress=progress.append
        )["status"]
        == "done"
    )
    assert "資料を読み取り専用で開く" in progress


@pytest.mark.parametrize(
    "failure,code",
    [
        (MemoryError("allocation failed"), "MEMORY_EXHAUSTED"),
        (PermissionError("denied"), "ACCESS_DENIED"),
    ],
)
def test_resource_failures_are_not_misreported_as_wrong_pptx(tmp_path, failure, code):
    source = tmp_path / "source.pptx"
    deck(source)

    def renderer(*args):
        raise failure

    app = create_app("test", tmp_path / "work", renderer)
    headers = {"X-Pawapo-Token": "test"}
    with TestClient(app, base_url="http://127.0.0.1") as client:
        response = client.post(
            "/api/jobs",
            headers=headers,
            files={
                side: ("source.pptx", source.read_bytes()) for side in ("left", "right")
            },
        )
        status = wait(client, response.json()["id"], headers)
        assert status["diagnostic"]["code"] == code
        assert status["diagnostic"]["actions"]
        assert "pptxとして保存" not in status["message"]


def test_worker_preserves_original_failure_stage_after_cleanup(tmp_path):
    from tests.test_render_com import office
    from pawapodiff.render import _render_powerpoint

    calls, state, dependencies = office(failures=["export", "close"])
    stages = []
    with pytest.raises(RenderError) as caught:
        _render_powerpoint(
            tmp_path / "private.pptx",
            tmp_path / "out",
            160,
            stage_callback=stages.append,
            **dependencies,
        )
    data = describe(caught.value)
    assert "PNG出力" in data["stage"]
    assert stages[-1] == "一時資料を閉じる"
    assert "close failed" in data["cleanup"]


def test_placeholder_picture_is_hashed_from_original_media(tmp_path):
    source = tmp_path / "source.pptx"
    picture = tmp_path / "image.png"
    Image.new("RGB", (12, 12), "red").save(picture)
    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[8])
    placeholder = next(
        shape for shape in slide.placeholders if hasattr(shape, "insert_picture")
    )
    placeholder.insert_picture(str(picture))
    p.save(source)
    result = engine.inspect(source)
    assert sum(bool(shape["image"]) for shape in result["slides"][0]) == 1


def test_diagnostics_do_not_need_writable_work_disk(tmp_path, monkeypatch):
    import pawapodiff.server as server

    app = create_app("test", tmp_path / "work")

    def unavailable(*args, **kwargs):
        raise PermissionError("Drive information unavailable")

    monkeypatch.setattr(server.shutil, "disk_usage", unavailable)
    with TestClient(app, base_url="http://127.0.0.1") as client:
        response = client.get("/api/diagnostics", headers={"X-Pawapo-Token": "test"})
        assert response.status_code == 200
        with ZipFile(BytesIO(response.content)) as z:
            assert "まだ比較ジョブがありません" in z.read(
                "diagnostic-summary.txt"
            ).decode("utf-8-sig")
        assert not list((tmp_path / "work").iterdir())


def test_missing_group_geometry_does_not_invent_child_movement(tmp_path):
    source = tmp_path / "source.pptx"
    p = deck(source)
    group = p.slides[0].shapes.add_group_shape()
    group.shapes.add_textbox(
        Inches(1), Inches(1), Inches(2), Inches(1)
    ).text = "グループ内の文章"
    group.left += Inches(1)
    group.width *= 2
    p.save(source)
    normal = engine.inspect(source)
    group._element.grpSpPr.remove(group._element.grpSpPr.xfrm)
    p.save(source)
    partial = engine.inspect(source)
    changes = engine.compare(normal, partial)["pages"][0]["changes"]
    assert not any(change.get("highlight") in ("layout", "mixed") for change in changes)


def test_unknown_inherited_format_is_not_called_a_confirmed_change(
    tmp_path, monkeypatch
):
    source = tmp_path / "source.pptx"
    deck(source)
    normal = engine.inspect(source)

    def unsupported(*args):
        raise ValueError("unsupported template metadata")

    monkeypatch.setattr(engine, "inherited_signature", unsupported)
    partial = engine.inspect(source)
    result = engine.compare(normal, partial)
    assert not result["pages"][0]["changes"]
    assert result["warnings"]
