import json
from pathlib import Path
import subprocess
import sys
import time

import pytest
from pawapodiff.render import _run_worker, RenderError


def command(tmp_path, body):
    state = tmp_path / "state.json"
    script = (
        "from pathlib import Path\nimport json,time,os,sys\nfrom pawapodiff.office_worker import write_state\nstate=Path("
        + repr(str(state))
        + ")\n"
        + body
    )
    return [sys.executable, "-c", script], state, tmp_path / "stderr.log"


def test_worker_success_protocol(tmp_path):
    args, state, stderr = command(
        tmp_path, "write_state(state,{'status':'done','images':['0001.png']})"
    )
    assert _run_worker(args, state, stderr, timeout=5)["images"] == ["0001.png"]


@pytest.mark.parametrize(
    "body,message",
    [
        (
            "write_state(state,{'status':'error','message':'Open failed: 0x80020003'});sys.exit(1)",
            "Open failed: 0x80020003",
        ),
        ("sys.stderr.write('native worker crash');sys.exit(7)", "終了コード7"),
        ("state.write_text('invalid JSON')", "正常に完了"),
    ],
)
def test_worker_failures_are_reported(tmp_path, body, message):
    args, state, stderr = command(tmp_path, body)
    with pytest.raises(RenderError, match=message):
        _run_worker(args, state, stderr, timeout=5)


def test_stuck_worker_is_terminated_without_killing_office(tmp_path, monkeypatch):
    args, state, stderr = command(
        tmp_path,
        "write_state(state,{'status':'running','stage':'資料を開く'});time.sleep(60)",
    )
    created = []
    popen = subprocess.Popen

    def capture(*args, **kwargs):
        child = popen(*args, **kwargs)
        created.append(child)
        return child

    monkeypatch.setattr(subprocess, "Popen", capture)
    started = time.monotonic()
    with pytest.raises(RenderError, match="応答しない"):
        _run_worker(args, state, stderr, timeout=1)
    assert time.monotonic() - started < 10
    assert len(created) == 1 and created[0].poll() is not None


def test_progress_resets_inactivity_timeout(tmp_path):
    body = "for i in range(10):\n write_state(state,{'status':'running','stage':str(i)});time.sleep(.3)\nwrite_state(state,{'status':'done'})"
    args, state, stderr = command(tmp_path, body)
    assert _run_worker(args, state, stderr, timeout=2)["status"] == "done"


def test_render_passes_only_unique_private_copies_to_office(tmp_path, monkeypatch):
    import pawapodiff.render as module
    from pptx import Presentation

    source = tmp_path / "original.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(source)
    original = source.read_bytes()
    copies = []

    def worker(args, state_path, stderr_path, **kwargs):
        request = json.loads(Path(args[-1]).read_text())
        private = Path(request["source"])
        assert private != source and private.read_bytes() == original
        assert private.parent != source.parent
        copies.append(private)
        out = Path(request["destination"]) / "0001.png"
        from PIL import Image

        Image.new("RGB", (160, 90)).save(out)
        return {"status": "done", "images": ["0001.png"]}

    monkeypatch.setattr(module.sys, "platform", "win32")
    monkeypatch.setattr(module, "_run_worker", worker)
    for index in range(2):
        assert len(module.render(source, tmp_path / f"out-{index}", 160)) == 1
    assert copies[0] != copies[1]
    assert all(not path.parent.exists() for path in copies)
    assert source.read_bytes() == original


def test_render_failure_preserves_original_and_cleans_staging(tmp_path, monkeypatch):
    import pawapodiff.render as module
    from pptx import Presentation

    source = tmp_path / "original.pptx"
    presentation = Presentation()
    presentation.save(source)
    original = source.read_bytes()
    copies = []

    def worker(args, *a, **kw):
        copies.append(Path(json.loads(Path(args[-1]).read_text())["source"]))
        raise RenderError("読み込み失敗 HRESULT 0x80020003")

    monkeypatch.setattr(module.sys, "platform", "win32")
    monkeypatch.setattr(module, "_run_worker", worker)
    with pytest.raises(RenderError, match="診断ログ"):
        module.render(source, tmp_path / "out")
    assert source.read_bytes() == original
    assert all(not path.parent.exists() for path in copies)
