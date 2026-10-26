from hashlib import sha256
from zipfile import ZipFile
from PIL import Image
from pptx import Presentation
from pawapodiff.engine import inspect


def test_analysis_never_bulk_reads_media_and_streams_image_hashes(
    tmp_path, monkeypatch
):
    image = tmp_path / "picture.png"
    Image.new("RGB", (20, 20), "red").save(image)
    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[6])
    slide.shapes.add_picture(str(image), 0, 0)
    source = tmp_path / "large-media.pptx"
    p.save(source)
    with ZipFile(source, "a") as archive:
        archive.writestr("ppt/media/unused-video.bin", b"x" * 1024**2)
    read, open_member = ZipFile.read, ZipFile.open
    sizes = []

    def guarded_read(self, name, *a, **kw):
        if str(self.filename) == str(source) and str(name).startswith("ppt/media/"):
            raise AssertionError("Original media must not be read in bulk")
        return read(self, name, *a, **kw)

    class Stream:
        def __init__(self, stream):
            self.stream = stream

        def __enter__(self):
            return self

        def __exit__(self, *a):
            self.stream.close()

        def read(self, size=-1):
            assert 0 < size <= 1024**2
            sizes.append(size)
            return self.stream.read(size)

    def guarded_open(self, name, *a, **kw):
        stream = open_member(self, name, *a, **kw)
        if str(self.filename) == str(source) and str(name).startswith("ppt/media/"):
            assert str(name).endswith(".png"), "Unused video must not be read"
            return Stream(stream)
        return stream

    monkeypatch.setattr(ZipFile, "read", guarded_read)
    monkeypatch.setattr(ZipFile, "open", guarded_open)
    result = inspect(source)
    assert result["slides"][0][0]["image"] == sha256(image.read_bytes()).hexdigest()
    assert sizes
