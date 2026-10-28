"""Actual PPTX movie structures; no codec, network, repair, or media playback."""

from hashlib import sha256
import json
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
import pytest

from pawapodiff.engine import inspect, validate, compare, InputError
from pawapodiff.media import MEDIA_TYPES, R, prepare_copy, parse, serialize


def video_deck(root, mode="embedded"):
    root.mkdir(exist_ok=True)
    movie, poster = root / "movie.mp4", root / "poster.png"
    movie.write_bytes(b"FAKE MEDIA: NEVER DECODE OR PLAY")
    Image.new("RGB", (40, 30), "red").save(poster)
    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = "動画付きの説明文章"
    slide.shapes.add_movie(
        str(movie),
        Inches(1),
        Inches(2),
        Inches(3),
        Inches(2),
        poster_frame_image=str(poster),
        mime_type="video/mp4",
    )
    original, changed = root / "original.pptx", root / "input.pptx"
    p.save(original)
    with ZipFile(original) as source, ZipFile(changed, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename.endswith(".mp4") and mode in (
                "missing",
                "external",
                "orphan",
            ):
                continue
            if info.filename.endswith(".png") and mode == "missing-poster":
                continue
            if info.filename == "ppt/slides/_rels/slide1.xml.rels":
                root_xml = parse(data, info.filename)
                for relationship in list(root_xml):
                    if relationship.get("Type", "").rsplit("/", 1)[-1] in MEDIA_TYPES:
                        if mode == "external":
                            relationship.set("TargetMode", "External")
                            relationship.set(
                                "Target",
                                "https://must-never-be-fetched.invalid/movie.mp4",
                            )
                        elif mode == "orphan":
                            root_xml.remove(relationship)
                data = serialize(root_xml)
            target.writestr(info, data)
    return changed


@pytest.mark.parametrize(
    "mode", ["embedded", "external", "missing", "missing-poster", "orphan"]
)
def test_video_variants_compare_as_static_without_repair_or_fetch(
    tmp_path, mode, monkeypatch
):
    import urllib.request

    source = video_deck(tmp_path, mode)
    original_hash = sha256(source.read_bytes()).hexdigest()

    def network_forbidden(*args, **kwargs):
        pytest.fail("Media must never be fetched")

    monkeypatch.setattr(urllib.request, "urlopen", network_forbidden)
    plan = validate(source)
    assert plan.notices and plan.notices[0]["page"] == 1
    destination = tmp_path / "private.pptx"
    prepare_copy(source, destination, plan)
    prepared = validate(destination)
    assert not prepared.replacements and not prepared.notices
    with ZipFile(destination) as package:
        for name in package.namelist():
            if name.endswith(".rels"):
                root = parse(package.read(name), name)
                assert not any(
                    r.get("Type", "").rsplit("/", 1)[-1] in MEDIA_TYPES for r in root
                )
        document = package.read("ppt/slides/slide1.xml")
        assert b"videoFile" not in document and b"ppaction://media" not in document
        assert b"<p:timing" not in document
        if mode != "orphan":
            assert not any(name.endswith(".mp4") for name in package.namelist())
    data = inspect(source)
    assert data["notices"] and not data["warnings"]
    assert len(data["slides"]) == 1
    assert any(shape["text"] == "動画付きの説明文章" for shape in data["slides"][0])
    pictures = [shape for shape in data["slides"][0] if shape["image"]]
    assert len(pictures) == (0 if mode == "missing-poster" else 1)
    result = compare(data, inspect(destination))
    assert result["notices"] and len(result["translations"]) == 1
    assert sha256(source.read_bytes()).hexdigest() == original_hash


def test_video_preparation_does_not_bulk_read_media(tmp_path, monkeypatch):
    source = video_deck(tmp_path)
    original_open, original_read = ZipFile.open, ZipFile.read

    def read(self, name, *args, **kwargs):
        if str(self.filename) == str(source):
            assert not str(name).endswith((".mp4", ".png"))
        return original_read(self, name, *args, **kwargs)

    def open_member(self, name, *args, **kwargs):
        label = name.filename if hasattr(name, "filename") else str(name)
        if str(self.filename) == str(source) and label.endswith(".mp4"):
            pytest.fail("Excluded movie bytes must not be read")
        return original_open(self, name, *args, **kwargs)

    monkeypatch.setattr(ZipFile, "read", read)
    monkeypatch.setattr(ZipFile, "open", open_member)
    plan = validate(source)
    prepare_copy(source, tmp_path / "private.pptx", plan)
    assert inspect(source)["notices"]


def test_non_video_broken_references_and_external_ole_still_fail(tmp_path):
    source = video_deck(tmp_path)
    with ZipFile(source) as original:
        members = {
            info.filename: original.read(info.filename) for info in original.infolist()
        }
    for mode in ("missing-slide", "external-ole"):
        target = tmp_path / f"{mode}.pptx"
        modified = dict(members)
        if mode == "missing-slide":
            del modified["ppt/slides/slide1.xml"]
        else:
            name = "ppt/slides/_rels/slide1.xml.rels"
            root = parse(modified[name], name)
            etree.SubElement(
                root,
                "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                Id="externalOle",
                Type=R + "/oleObject",
                Target="https://must-never-be-fetched.invalid/data.xlsx",
                TargetMode="External",
            )
            modified[name] = serialize(root)
        with ZipFile(target, "w") as archive:
            for name, data in modified.items():
                archive.writestr(name, data)
        with pytest.raises(InputError) as caught:
            inspect(target)
        assert caught.value.diagnostic["code"] == (
            "MISSING_PART" if mode == "missing-slide" else "EXTERNAL_REFERENCE"
        )


def test_renderer_passes_only_static_private_copy_to_office(tmp_path, monkeypatch):
    import pawapodiff.render as module

    source = video_deck(tmp_path, "external")
    original = source.read_bytes()

    def worker(args, *rest, **kwargs):
        request = json.loads(Path(args[-1]).read_text())
        private = Path(request["source"])
        assert private != source and private.read_bytes() != original
        assert not validate(private).notices
        output = Path(request["destination"]) / "0001.png"
        Image.new("RGB", (160, 90)).save(output)
        return {"status": "done", "images": [output.name]}

    monkeypatch.setattr(module.sys, "platform", "win32")
    monkeypatch.setattr(module, "_run_worker", worker)
    assert module.render(source, tmp_path / "images", 160)
    assert source.read_bytes() == original


def test_audio_transition_is_removed_without_dangling_sound_element(tmp_path):
    source = video_deck(tmp_path)
    target = tmp_path / "audio.pptx"
    with ZipFile(source) as original, ZipFile(target, "w") as output:
        for info in original.infolist():
            data = original.read(info.filename)
            if info.filename == "ppt/slides/slide1.xml":
                document = parse(data, info.filename)
                from pawapodiff.media import P, NS

                for timing in document.findall("p:timing", NS):
                    document.remove(timing)
                transition = etree.SubElement(document, "{" + P + "}transition")
                action = etree.SubElement(transition, "{" + P + "}sndAc")
                start = etree.SubElement(action, "{" + P + "}stSnd")
                etree.SubElement(
                    start, "{" + P + "}snd", {"{" + R + "}embed": "audioExternal"}
                )
                data = serialize(document)
            elif info.filename == "ppt/slides/_rels/slide1.xml.rels":
                relationships = parse(data, info.filename)
                etree.SubElement(
                    relationships,
                    "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                    Id="audioExternal",
                    Type=R + "/audio",
                    Target="https://never-fetch.invalid/sound.wav",
                    TargetMode="External",
                )
                data = serialize(relationships)
            output.writestr(info, data)
    prepared = tmp_path / "static.pptx"
    prepare_copy(target, prepared, validate(target))
    with ZipFile(prepared) as package:
        assert b"sndAc" not in package.read("ppt/slides/slide1.xml")
    assert not validate(prepared).notices


def test_media_geometry_poster_and_non_media_shapes_are_preserved(tmp_path):
    source = video_deck(tmp_path)
    original = Presentation(source)
    document = inspect(source)
    assert len(document["slides"][0]) == len(original.slides[0].shapes)
    for actual, expected in zip(document["slides"][0], original.slides[0].shapes):
        assert actual["id"] == str(expected.shape_id)
        assert actual["box"] == [
            getattr(expected, attribute) / 12700
            for attribute in ("left", "top", "width", "height")
        ]
    assert (
        document["slides"][0][1]["image"]
        == sha256((tmp_path / "poster.png").read_bytes()).hexdigest()
    )


def test_video_api_job_completes_and_records_notices_in_report_and_diagnostics(
    tmp_path,
):
    from io import BytesIO
    from fastapi.testclient import TestClient
    from pawapodiff.server import create_app
    from tests.test_loading_failures import wait

    left = video_deck(tmp_path / "left", "external")
    right = video_deck(tmp_path / "right", "missing-poster")

    def renderer(path, destination):
        destination.mkdir()
        assert inspect(path)["notices"]
        target = destination / "0001.png"
        Image.new("RGB", (160, 90)).save(target)
        return [target]

    app = create_app("test", tmp_path / "work", renderer)
    headers = {"X-Pawapo-Token": "test"}
    with TestClient(app, base_url="http://127.0.0.1") as client:
        response = client.post(
            "/api/jobs",
            headers=headers,
            files={
                side: ("video.pptx", path.read_bytes())
                for side, path in (("left", left), ("right", right))
            },
        )
        key = response.json()["id"]
        status = wait(client, key, headers)
        assert status["status"] == "done", status
        assert len(status["result"]["notices"]) == 2
        html = client.get(f"/api/jobs/{key}/report", headers=headers).text
        assert "動画・音声は静止画として比較します" in html
        assert "must-never-be-fetched" not in html
        with ZipFile(
            BytesIO(client.get("/api/diagnostics", headers=headers).content)
        ) as package:
            snapshot = json.loads(package.read("diagnostic-summary.json"))
            assert snapshot["jobs"][key]["files"]["left"]["media_notices"]
