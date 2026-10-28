"""User-visible structural changes, granular styles, and placeholder identity."""

from copy import deepcopy
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import pytest

from pawapodiff.engine import inspect, compare
from pawapodiff.report import report


def text_deck(path):
    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text = "原文の文章"
    font = box.text_frame.paragraphs[0].runs[0].font
    font.name, font.size, font.bold, font.italic = "Arial", Pt(20), False, False
    font.color.rgb = RGBColor(0, 0, 0)
    p.save(path)
    return p


def fields(result):
    return [
        field
        for page in result["pages"]
        for change in page["changes"]
        for field in change["fields"]
    ]


@pytest.mark.parametrize(
    "attribute,value,label,before,after",
    [
        ("bold", True, "太字（ボールド）", "オフ", "オン"),
        ("size", Pt(24), "文字サイズ", "20 pt", "24 pt"),
        ("name", "Calibri", "フォント名", "Arial", "Calibri"),
        ("italic", True, "斜体", "オフ", "オン"),
        ("color", RGBColor(255, 0, 0), "文字色", "#000000", "#FF0000"),
    ],
)
def test_each_font_attribute_is_named_with_before_and_after(
    tmp_path, attribute, value, label, before, after
):
    source = tmp_path / "source.pptx"
    p = text_deck(source)
    original = inspect(source)
    run = p.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    run.text = "Translated content"
    if attribute == "color":
        run.font.color.rgb = value
    else:
        setattr(run.font, attribute, value)
    p.save(source)
    result = compare(original, inspect(source))
    assert len(fields(result)) == 1, fields(result)
    detail = fields(result)[0]
    assert detail.startswith(label + ":") and before in detail and after in detail
    assert result["pages"][0]["changes"][0]["highlight"] == "format"
    assert "文字内容" not in detail and "文字書式・配置" not in detail


def test_translation_only_and_empty_translation_do_not_create_change_list_entries(
    tmp_path,
):
    source = tmp_path / "source.pptx"
    p = text_deck(source)
    original = inspect(source)
    box = p.slides[0].shapes[0]
    box.text_frame.paragraphs[0].runs[0].text = "Translated"
    p.save(source)
    translated = compare(original, inspect(source))
    assert translated["pages"][0]["changes"] == []
    assert translated["translations"][0]["target"] == "Translated"
    box.text = ""
    p.save(source)
    missing = compare(original, inspect(source))
    assert missing["pages"][0]["changes"] == []
    assert "訳抜け候補" in missing["translations"][0]["notes"]


def test_partial_bold_is_named_without_unrelated_font_attributes(tmp_path):
    source = tmp_path / "source.pptx"
    p = text_deck(source)
    paragraph = p.slides[0].shapes[0].text_frame.paragraphs[0]
    extra = paragraph.add_run()
    extra.text = "続き"
    extra.font.name, extra.font.size, extra.font.bold, extra.font.italic = (
        "Arial",
        Pt(20),
        False,
        False,
    )
    extra.font.color.rgb = RGBColor(0, 0, 0)
    p.save(source)
    original = inspect(source)
    extra.font.bold = True
    p.save(source)
    details = fields(compare(original, inspect(source)))
    assert len(details) == 1
    assert details[0].startswith("太字（ボールド）") and "部分書式" in details[0]


def test_paragraph_defaults_alignment_and_autofit_are_described(tmp_path):
    source = tmp_path / "source.pptx"
    p = text_deck(source)
    box = p.slides[0].shapes[0]
    paragraph = box.text_frame.paragraphs[0]
    paragraph.runs[0].font.bold = None
    p.save(source)
    original = inspect(source)
    paragraph.font.bold = True
    paragraph.alignment = PP_ALIGN.CENTER
    box.text_frame.margin_left = Pt(12)
    p.save(source)
    details = fields(compare(original, inspect(source)))
    assert any(detail.startswith("太字（ボールド）") for detail in details)
    assert any(
        detail.startswith("段落の配置") and "中央揃え" in detail for detail in details
    )
    assert any(
        detail.startswith("文字ボックスの左余白") and "12 pt" in detail
        for detail in details
    )


def placeholder_deck(path, with_text=False):
    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[1])
    if with_text:
        slide.placeholders[1].text = "削除されるプレースホルダーの原文"
    p.save(path)
    return p


def test_empty_placeholder_addition_and_deletion_are_highlighted_even_when_id_reused(
    tmp_path,
):
    source = tmp_path / "source.pptx"
    p = placeholder_deck(source)
    original = inspect(source)
    slide = p.slides[0]
    old = slide.placeholders[1]
    dimensions = [old.left, old.top, old.width, old.height]
    new = deepcopy(old._element)
    slide.shapes._spTree.remove(old._element)
    slide.shapes._spTree.insert_element_before(new, "p:extLst")
    added = slide.shapes[-1]
    for key, value in zip(("left", "top", "width", "height"), dimensions):
        setattr(added, key, value)
    added._element.xpath(".//p:ph")[0].set("idx", "9")
    p.save(source)
    result = compare(original, inspect(source))
    changes = result["pages"][0]["changes"]
    assert {field for change in changes for field in change["fields"]} == {
        "プレースホルダー追加",
        "プレースホルダー削除",
    }
    assert len(changes) == 2 and all(
        change["highlight"] == "other" for change in changes
    )
    assert any(change["left"] and not change["right"] for change in changes)
    assert any(change["right"] and not change["left"] for change in changes)
    assert not any(change["uncertain"] for change in changes)


def test_placeholder_identity_survives_shape_id_changes(tmp_path):
    source = tmp_path / "source.pptx"
    p = placeholder_deck(source, with_text=True)
    original = inspect(source)
    placeholder = p.slides[0].placeholders[1]
    placeholder._element.xpath(".//p:cNvPr")[0].set("id", "99")
    placeholder.text_frame.paragraphs[0].runs[0].text = "Translated"
    p.save(source)
    result = compare(original, inspect(source))
    assert not result["pages"][0]["changes"]
    assert result["translations"][0]["target"] == "Translated"
    assert not result["translations"][0]["uncertain"]


def test_deleted_placeholder_is_not_a_missing_translation(tmp_path):
    source = tmp_path / "source.pptx"
    p = placeholder_deck(source, with_text=True)
    original = inspect(source)
    removed = p.slides[0].placeholders[1]
    p.slides[0].shapes._spTree.remove(removed._element)
    p.save(source)
    result = compare(original, inspect(source))
    assert fields(result) == ["プレースホルダー削除"]
    assert not result["issues"]
    row = result["translations"][0]
    assert "訳抜け候補" not in row["notes"]
    assert any("プレースホルダー削除" in note for note in row["notes"])


def test_existing_placeholder_with_empty_translation_is_still_a_candidate(tmp_path):
    source = tmp_path / "source.pptx"
    p = placeholder_deck(source, with_text=True)
    original = inspect(source)
    p.slides[0].placeholders[1].text = ""
    p.save(source)
    result = compare(original, inspect(source))
    assert not result["pages"][0]["changes"]
    assert "訳抜け候補" in result["translations"][0]["notes"]


def test_master_changes_are_not_listed_and_html_has_only_bilingual_table(tmp_path):
    source = tmp_path / "source.pptx"
    p = text_deck(source)
    original = inspect(source)
    p.slide_master._element.set("preserve", "1")
    p.save(source)
    result = compare(original, inspect(source))
    assert not result["pages"][0]["changes"]
    image = tmp_path / "image.png"
    Image.new("RGB", (160, 90)).save(image)
    html = report(result, {"left": [image], "right": [image]})
    assert "<h2>日英比較表</h2>" in html
    assert "翻訳漏れ・表記揺れチェック" not in html
    assert "文字内容（翻訳を含む）" not in html
    assert "テーマ・マスター・レイアウト" not in html


def test_added_placeholder_is_not_an_untranslated_text_candidate(tmp_path):
    source = tmp_path / "source.pptx"
    p = placeholder_deck(source)
    slide = p.slides[0]
    body = slide.placeholders[1]
    element = deepcopy(body._element)
    slide.shapes._spTree.remove(body._element)
    p.save(source)
    original = inspect(source)
    slide.shapes._spTree.insert_element_before(element, "p:extLst")
    slide.placeholders[1].text = "追加された文章"
    p.save(source)
    result = compare(original, inspect(source))
    assert fields(result) == ["プレースホルダー追加"]
    assert not result["issues"]


def test_empty_placeholder_default_font_changes_are_named(tmp_path):
    source = tmp_path / "source.pptx"
    p = placeholder_deck(source)
    paragraph = p.slides[0].placeholders[1].text_frame.paragraphs[0]
    paragraph.font.size = Pt(20)
    p.save(source)
    original = inspect(source)
    paragraph.font.size = Pt(24)
    p.save(source)
    assert fields(compare(original, inspect(source))) == [
        "文字サイズ: 20 pt → 24 pt（段落の既定書式）"
    ]


def test_picture_inserted_in_existing_placeholder_is_an_image_change(tmp_path):
    source = tmp_path / "source.pptx"
    image = tmp_path / "image.png"
    Image.new("RGB", (40, 30), "red").save(image)
    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[8])
    p.save(source)
    original = inspect(source)
    placeholder = next(
        item for item in slide.placeholders if hasattr(item, "insert_picture")
    )
    placeholder.insert_picture(str(image))
    p.save(source)
    result = compare(original, inspect(source))
    assert fields(result) == ["画像"]
    assert result["pages"][0]["changes"][0]["highlight"] == "other"


def test_east_asian_font_and_autofit_scaling_are_reported(tmp_path):
    from pptx.oxml.xmlchemy import OxmlElement

    source = tmp_path / "source.pptx"
    p = text_deck(source)
    box = p.slides[0].shapes[0]
    rpr = box.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr()
    ea = OxmlElement("a:ea")
    ea.set("typeface", "游ゴシック")
    rpr.append(ea)
    body = box.text_frame._txBody.bodyPr
    for node in list(body):
        body.remove(node)
    autofit = OxmlElement("a:normAutofit")
    autofit.set("fontScale", "100000")
    body.append(autofit)
    p.save(source)
    original = inspect(source)
    ea.set("typeface", "メイリオ")
    autofit.set("fontScale", "80000")
    p.save(source)
    details = fields(compare(original, inspect(source)))
    assert any(
        detail.startswith("フォント名（東アジア文字）")
        and "游ゴシック" in detail
        and "メイリオ" in detail
        for detail in details
    )
    assert any(
        detail.startswith("自動調整の文字倍率") and "100%" in detail and "80%" in detail
        for detail in details
    )
