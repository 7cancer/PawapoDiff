"""Run on Windows with desktop PowerPoint. Produces no repository artifacts."""

import sys
from pathlib import Path
import tempfile
from zipfile import ZipFile
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from pptx import Presentation
from pptx.util import Inches, Pt
from pawapodiff.engine import inspect, compare
from pawapodiff.render import render
from pawapodiff.report import report
from pawapodiff.media import MEDIA_TYPES, parse, serialize


def main():
    if sys.platform != "win32":
        raise SystemExit(
            "This acceptance check requires Windows and Microsoft PowerPoint."
        )
    with tempfile.TemporaryDirectory(prefix="PawapoDiff-acceptance-") as directory:
        root = Path(directory)
        movie, poster = root / "unused-video.mp4", root / "poster.png"
        # The payload is intentionally not a playable video: it must be removed
        # from the private copy before PowerPoint is allowed to see the document.
        movie.write_bytes(b"STATIC COMPARISON TEST: DO NOT PLAY")
        Image.new("RGB", (40, 30), "steelblue").save(poster)
        for translated in (False, True):
            p = Presentation()
            p.slide_width = Inches(12)
            p.slide_height = Inches(6.75)
            for i in range(2):
                s = p.slides.add_slide(p.slide_layouts[6])
                t = s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
                t.text = ["Sales", "Revenue"][i] if translated else "売上高"
                t.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
                if translated:
                    t.left += Pt(3)
            p.slides[0].shapes.add_movie(
                str(movie),
                Inches(6),
                Inches(3),
                Inches(3),
                Inches(2),
                poster_frame_image=str(poster),
                mime_type="video/mp4",
            )
            generated = root / "generated.pptx"
            p.save(generated)
            with (
                ZipFile(generated) as original,
                ZipFile(
                    root / ("right.pptx" if translated else "left.pptx"), "w"
                ) as output,
            ):
                for info in original.infolist():
                    if info.filename.endswith(".mp4"):
                        continue
                    data = original.read(info.filename)
                    if info.filename.endswith(".rels"):
                        relationships = parse(data, info.filename)
                        for rel in relationships:
                            if rel.get("Type", "").rsplit("/", 1)[-1] in MEDIA_TYPES:
                                rel.set(
                                    "Target",
                                    "../media/missing.mp4"
                                    if translated
                                    else "https://never-fetch.invalid/movie.mp4",
                                )
                                if not translated:
                                    rel.set("TargetMode", "External")
                        data = serialize(relationships)
                    output.writestr(info, data)
        result = compare(inspect(root / "left.pptx"), inspect(root / "right.pptx"))
        images = {
            side: render(root / f"{side}.pptx", root / side)
            for side in ("left", "right")
        }
        assert all(len(paths) == 2 for paths in images.values())
        assert any(i["kind"] == "表記揺れ候補" for i in result["issues"])
        assert len(result["notices"]) == 2
        assert len(result["translations"]) == 1
        assert len(result["translations"][0]["entries"]) == 2
        assert all(row["inconsistent"] for row in result["translations"])
        html = report(result, images)
        assert html.count("data:image/png;base64,") == 4
        print(
            "PASS: PowerPoint static rendering with external/missing video references, grouped translation table, embedded HTML."
        )
        print(
            "Packaged EXE and existing PowerPoint process behavior still need the README manual checks."
        )


if __name__ == "__main__":
    main()
