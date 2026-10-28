"""GIF and optional-image references must not block usable static slides."""

from hashlib import sha256
from pathlib import Path
from urllib.parse import quote
from zipfile import ZipFile

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
import pytest

from pawapodiff.engine import inspect, validate, InputError
from pawapodiff.media import A, R, NS, parse, serialize, prepare_copy

REL = "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"


def gif_deck(root, mode="normal"):
    root.mkdir(exist_ok=True)
    gif = root / "animated.gif"
    frames = [Image.new("RGB", (40, 30), color) for color in ("red", "blue")]
    frames[0].save(gif, save_all=True, append_images=frames[1:], duration=100, loop=0)
    png = root / "cached.png"
    Image.new("RGB", (40, 30), "green").save(png)
    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = "GIFがある説明文章"
    slide.shapes.add_picture(str(gif), Inches(1), Inches(2), Inches(3), Inches(2))
    original, source = root / "original.pptx", root / "input.pptx"
    p.save(original)
    with ZipFile(original) as package:
        members = {name: package.read(name) for name in package.namelist()}
    owner, rels = "ppt/slides/slide1.xml", "ppt/slides/_rels/slide1.xml.rels"
    document, relationships = parse(members[owner], owner), parse(members[rels], rels)
    part = "ppt/media/image1.gif"
    if mode in ("missing", "cached"):
        del members[part]
    if mode == "undefined":
        document.find(".//a:blip", NS).set("{" + R + "}embed", "undefinedGif")
    if mode in ("unused", "optional", "optional-undefined", "cached"):
        if mode != "optional-undefined":
            etree.SubElement(
                relationships,
                REL,
                Id="optionalGif",
                Type=R + ("/hdphoto" if mode == "optional" else "/image"),
                Target="../media/cached.png"
                if mode == "cached"
                else "../media/missing.gif",
            )
        if mode != "unused":
            blip = document.find(".//a:blip", NS)
            extension = etree.SubElement(
                etree.SubElement(blip, "{" + A + "}extLst"),
                "{" + A + "}ext",
                uri="{TEST-IMAGE-PROCESSING}",
            )
            etree.SubElement(
                extension,
                "{http://schemas.microsoft.com/office/drawing/2010/main}imgLayer",
                {"{" + R + "}embed": "optionalGif"},
            )
        if mode == "cached":
            members["ppt/media/cached.png"] = png.read_bytes()
            types = parse(members["[Content_Types].xml"], "[Content_Types].xml")
            etree.SubElement(
                types,
                "{http://schemas.openxmlformats.org/package/2006/content-types}Default",
                Extension="png",
                ContentType="image/png",
            )
            members["[Content_Types].xml"] = serialize(types)
    members[owner], members[rels] = serialize(document), serialize(relationships)
    with ZipFile(source, "w") as package:
        for name, data in members.items():
            package.writestr(name, data)
    return source


@pytest.mark.parametrize(
    "mode",
    [
        "normal",
        "unused",
        "optional",
        "optional-undefined",
        "missing",
        "undefined",
        "cached",
    ],
)
def test_gif_image_references_keep_other_content_and_original(
    tmp_path, mode, monkeypatch
):
    import urllib.request

    source = gif_deck(tmp_path, mode)
    original = source.read_bytes()

    def never_fetch(*args, **kwargs):
        pytest.fail("Image references must never be fetched")

    monkeypatch.setattr(urllib.request, "urlopen", never_fetch)
    data = inspect(source)
    assert len(data["slides"]) == 1 and not data["warnings"]
    assert any(shape["text"] == "GIFがある説明文章" for shape in data["slides"][0])
    pictures = [shape for shape in data["slides"][0] if shape["image"]]
    assert len(pictures) == (0 if mode in ("missing", "undefined") else 1)
    if pictures:
        expected = (
            tmp_path / ("cached.png" if mode == "cached" else "animated.gif")
        ).read_bytes()
        assert pictures[0]["image"] == sha256(expected).hexdigest()
    assert bool(data["notices"]) == (mode != "normal")
    static = tmp_path / "static.pptx"
    prepare_copy(source, static, validate(source))
    assert not validate(static).replacements
    assert not inspect(static)["notices"]
    assert source.read_bytes() == original
    assert data["slides"][0][1]["box"] == [72, 144, 216, 144]


@pytest.mark.parametrize("filename", ["image 1.gif", "image%201.gif", "猫.gif"])
def test_gif_uri_names_resolve_existing_assets_without_false_missing_part(
    tmp_path, filename
):
    original = gif_deck(tmp_path)
    source = tmp_path / "escaped.pptx"
    with ZipFile(original) as archive, ZipFile(source, "w") as output:
        for info in archive.infolist():
            name, data = info.filename, archive.read(info.filename)
            if name == "ppt/media/image1.gif":
                name = "ppt/media/" + filename
            if name == "ppt/slides/_rels/slide1.xml.rels":
                root = parse(data, name)
                for relationship in root:
                    if relationship.get("Type", "").endswith("/image"):
                        relationship.set(
                            "Target", "../media/" + quote(filename, safe="")
                        )
                data = serialize(root)
            output.writestr(name, data)
    result = inspect(source)
    assert (
        result["slides"][0][1]["image"]
        == sha256((tmp_path / "animated.gif").read_bytes()).hexdigest()
    )
    private = tmp_path / "static.pptx"
    prepare_copy(source, private, validate(source))
    assert not validate(private).replacements


def test_missing_mandatory_slide_is_still_an_error(tmp_path):
    source = gif_deck(tmp_path)
    broken = tmp_path / "broken.pptx"
    with ZipFile(source) as original, ZipFile(broken, "w") as output:
        for info in original.infolist():
            if info.filename != "ppt/slides/slide1.xml":
                output.writestr(info, original.read(info.filename))
    with pytest.raises(InputError) as caught:
        inspect(broken)
    assert caught.value.diagnostic["code"] == "MISSING_PART"


def test_static_gif_copy_passed_to_office_has_no_broken_references(
    tmp_path, monkeypatch
):
    import json
    import pawapodiff.render as module

    source = gif_deck(tmp_path, "optional")
    original = source.read_bytes()

    def worker(args, *rest, **kwargs):
        request = json.loads(Path(args[-1]).read_text())
        private = Path(request["source"])
        assert private != source and not validate(private).replacements
        assert inspect(private)["slides"][0][1]["image"]
        output = Path(request["destination"]) / "0001.png"
        Image.new("RGB", (160, 90)).save(output)
        return {"status": "done", "images": [output.name]}

    monkeypatch.setattr(module.sys, "platform", "win32")
    monkeypatch.setattr(module, "_run_worker", worker)
    assert module.render(source, tmp_path / "images", 160)
    assert source.read_bytes() == original


def test_gif_job_reaches_done_and_notices_appear_in_offline_report_and_diagnostics(
    tmp_path,
):
    from io import BytesIO
    import json
    from fastapi.testclient import TestClient
    from pawapodiff.server import create_app
    from tests.test_loading_failures import wait

    source = gif_deck(tmp_path, "missing")

    def renderer(path, destination):
        destination.mkdir()
        target = destination / "0001.png"
        Image.new("RGB", (160, 90)).save(target)
        return [target]

    headers = {"X-Pawapo-Token": "test"}
    with TestClient(
        create_app("test", tmp_path / "work", renderer), base_url="http://127.0.0.1"
    ) as client:
        response = client.post(
            "/api/jobs",
            headers=headers,
            files={
                side: ("gif.pptx", source.read_bytes()) for side in ("left", "right")
            },
        )
        key = response.json()["id"]
        status = wait(client, key, headers)
        assert status["status"] == "done", status
        assert len(status["result"]["translations"]) == 1
        assert all(
            notice["code"] == "IMAGE_MISSING" for notice in status["result"]["notices"]
        )
        html = client.get("/api/jobs/" + key + "/report", headers=headers).text
        assert "画像・メディアの確認事項" in html and "代替枠" in html
        with ZipFile(
            BytesIO(client.get("/api/diagnostics", headers=headers).content)
        ) as package:
            data = json.loads(package.read("diagnostic-summary.json"))
            notice = data["jobs"][key]["files"]["left"]["media_notices"][0]
            assert notice["part"] == "ppt/slides/slide1.xml"
            assert notice["missing_parts"] == ["ppt/media/image1.gif"]


def test_gif_fallback_can_share_a_slide_with_video(tmp_path):
    from tests.test_media import video_deck

    source = video_deck(tmp_path)
    broken = tmp_path / "combined.pptx"
    with ZipFile(source) as original, ZipFile(broken, "w") as output:
        for info in original.infolist():
            data = original.read(info.filename)
            if info.filename == "ppt/slides/_rels/slide1.xml.rels":
                root = parse(data, info.filename)
                etree.SubElement(
                    root,
                    REL,
                    Id="unusedGif",
                    Type=R + "/image",
                    Target="../media/missing.gif",
                )
                data = serialize(root)
            output.writestr(info, data)
    data = inspect(broken)
    assert {"MEDIA_IGNORED", "IMAGE_FALLBACK"} == {
        notice["code"] for notice in data["notices"]
    }
    target = tmp_path / "static.pptx"
    prepare_copy(broken, target, validate(broken))
    assert not validate(target).replacements


def test_missing_image_fill_preserves_text_and_shape_geometry(tmp_path):
    original = gif_deck(tmp_path)
    source = tmp_path / "fill.pptx"
    with ZipFile(original) as package, ZipFile(source, "w") as output:
        for info in package.infolist():
            data = package.read(info.filename)
            if info.filename == "ppt/slides/slide1.xml":
                root = parse(data, info.filename)
                properties = root.find(".//p:sp/p:spPr", NS)
                fill = etree.SubElement(properties, "{" + A + "}blipFill")
                etree.SubElement(
                    fill, "{" + A + "}blip", {"{" + R + "}embed": "missingFill"}
                )
                etree.SubElement(
                    etree.SubElement(fill, "{" + A + "}stretch"), "{" + A + "}fillRect"
                )
                data = serialize(root)
            if info.filename == "ppt/slides/_rels/slide1.xml.rels":
                root = parse(data, info.filename)
                etree.SubElement(
                    root,
                    REL,
                    Id="missingFill",
                    Type=R + "/image",
                    Target="../media/missing-background.gif",
                )
                data = serialize(root)
            output.writestr(info, data)
    document = inspect(source)
    assert document["slides"][0][0]["text"] == "GIFがある説明文章"
    assert document["notices"][0]["code"] == "IMAGE_MISSING"
    assert document["notices"][0]["shape_id"] == "2"
    private = tmp_path / "static.pptx"
    prepare_copy(source, private, validate(source))
    assert not validate(private).replacements
