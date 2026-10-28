from pathlib import Path
import pytest
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
from pawapodiff.engine import inspect, compare, InputError
from pawapodiff.render import render, RenderError
from pawapodiff.report import report


def fixture(path, translated=False):
    p = Presentation()
    p.slide_width = Inches(12)
    p.slide_height = Inches(6.75)
    for i in range(2):
        s = p.slides.add_slide(p.slide_layouts[6])
        t = s.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        t.text = "売上高" if not translated else ["Sales", "Revenue"][i]
        t.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
        s.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(2), Inches(3), Inches(2), Inches(1)
        )
    p.save(path)
    return p


def test_translation_does_not_become_formatting(tmp_path):
    a, b = tmp_path / "a.pptx", tmp_path / "b.pptx"
    fixture(a)
    fixture(b, True)
    r = compare(inspect(a), inspect(b))
    assert all(c["kind"] == "text" for page in r["pages"] for c in page["changes"])
    assert [i["kind"] for i in r["issues"]] == ["表記揺れ候補"]
    assert {e["page"] for e in r["issues"][0]["entries"]} == {1, 2}


def test_geometry_partial_bold_font_color_rotation_image(tmp_path):
    image = tmp_path / "pic.png"
    Image.new("RGB", (10, 10), "red").save(image)
    a, b = tmp_path / "a.pptx", tmp_path / "b.pptx"
    p = fixture(a)
    p.slides[0].shapes.add_picture(str(image), 0, 0)
    p.save(a)
    p = Presentation(a)
    text = p.slides[0].shapes[0]
    text.left += Pt(2)
    text.width += Pt(3)
    text.rotation = 8
    run = text.text_frame.paragraphs[0].add_run()
    run.text = "追加"
    run.font.bold = True
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor(255, 0, 0)
    Image.new("RGB", (10, 10), "blue").save(image)
    p.slides[0].shapes._spTree.remove(p.slides[0].shapes[-1]._element)
    p.slides[0].shapes.add_picture(str(image), 0, 0)
    p.save(b)
    r = compare(inspect(a), inspect(b))
    fields = [f for c in r["pages"][0]["changes"] for f in c["fields"]]
    assert "横位置: +2.00 pt" in fields and "幅: +3.00 pt" in fields
    assert "回転" in fields and "文字書式・配置" in fields and "画像" in fields
    r = compare(inspect(a), inspect(b), 4)
    assert not any("横位置" in f for c in r["pages"][0]["changes"] for f in c["fields"])


def test_groups_tables_missing_japanese(tmp_path):
    a, b = tmp_path / "a.pptx", tmp_path / "b.pptx"
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[6])
    g = s.shapes.add_group_shape()
    sh = g.shapes.add_textbox(Inches(2), Inches(1), Inches(2), Inches(1))
    sh.text = "原文"
    table = s.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(2)).table
    table.cell(0, 0).text = "日本語"
    table.cell(1, 0).text = "翻訳"
    p.save(a)
    p = Presentation(a)
    p.slides[0].shapes[0].shapes[0].text = ""
    p.slides[0].shapes[0].left += Pt(5)
    p.slides[0].shapes[1].table.cell(1, 0).text = "Translation"
    p.save(b)
    r = compare(inspect(a), inspect(b))
    assert {"訳抜け候補", "日本語の残存候補"} <= {i["kind"] for i in r["issues"]}
    assert any(
        c["left"] and "/" in c["left"]["id"] and c["kind"] == "structure"
        for c in r["pages"][0]["changes"]
    )


def test_reject_dimensions_pages_corruption(tmp_path):
    a, b = tmp_path / "a.pptx", tmp_path / "b.pptx"
    fixture(a)
    p = fixture(b)
    p.slide_width += Pt(1)
    p.save(b)
    with pytest.raises(InputError, match="寸法"):
        compare(inspect(a), inspect(b))
    p.slides.add_slide(p.slide_layouts[6])
    p.save(b)
    with pytest.raises(InputError, match="ページ数"):
        compare(inspect(a), inspect(b))
    b.write_bytes(b"bad")
    with pytest.raises(InputError, match="暗号化"):
        inspect(b)


def test_embedded_report_escapes_untrusted_text(tmp_path):
    image = tmp_path / "slide.png"
    Image.new("RGB", (160, 90), "white").save(image)
    content = report(
        {
            "width": 100,
            "height": 100,
            "pages": [{"page": 1, "changes": []}],
            "issues": [
                {"kind": "訳抜け", "source": "<script>", "page": 1, "target": ""}
            ],
            "translations": [
                {
                    "page": 1,
                    "shape": "<img>",
                    "source": "<script>",
                    "target": "<b>",
                    "inconsistent": True,
                    "notes": ["表記揺らぎの可能性あり"],
                }
            ],
        },
        {"left": [image], "right": [image]},
    )
    assert "data:image/png;base64," in content and "&lt;script&gt;" in content
    assert content.count("<script>") == 1  # Only our fixed offline interaction script.
    assert "https://" not in content
    assert "日英比較表" in content and "表記揺らぎの可能性あり" in content


@pytest.mark.skipif(
    __import__("sys").platform == "win32",
    reason="Linux-specific unsupported platform check",
)
def test_linux_render_error(tmp_path):
    a = tmp_path / "a.pptx"
    fixture(a)
    with pytest.raises(RenderError, match="Windows"):
        render(a, tmp_path / "render")


def test_inherited_format_and_paragraph_defaults(tmp_path):
    a, b = tmp_path / "a.pptx", tmp_path / "b.pptx"
    fixture(a)
    p = Presentation(a)
    p.slides[0].shapes[0].text_frame.paragraphs[0].font.bold = True
    p.save(b)
    r = compare(inspect(a), inspect(b))
    assert any("文字書式・配置" in c["fields"] for c in r["pages"][0]["changes"])
    p = Presentation(a)
    from pptx.oxml.xmlchemy import OxmlElement

    p.slide_master._element.append(OxmlElement("p:clrMap"))
    p.save(b)
    r = compare(inspect(a), inspect(b))
    assert r["pages"][0]["changes"][0]["left"]["id"] == "inherited"


def test_fallback_matching_and_unmatched_elements(tmp_path):
    a, b = tmp_path / "a.pptx", tmp_path / "b.pptx"
    fixture(a)
    fixture(b, True)
    left, right = inspect(a), inspect(b)
    right["slides"][0][0]["id"] = "new-id"
    r = compare(left, right)
    assert any(c["uncertain"] for c in r["pages"][0]["changes"])
    right["slides"][0] = []
    r = compare(left, right)
    assert all(c["fields"] == ["要素の削除"] for c in r["pages"][0]["changes"])


def test_external_relationship_is_rejected(tmp_path):
    from zipfile import ZipFile

    a, b = tmp_path / "a.pptx", tmp_path / "b.pptx"
    fixture(a)
    with ZipFile(a) as source, ZipFile(b, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "ppt/_rels/presentation.xml.rels":
                data = data.replace(
                    b"</Relationships>",
                    b'<Relationship Id="external" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject" TargetMode="External" Target="file:///not-opened"/></Relationships>',
                )
            target.writestr(item, data)
    with pytest.raises(InputError, match="外部参照"):
        inspect(b)
